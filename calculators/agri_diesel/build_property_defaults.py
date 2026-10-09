"""
Build property_defaults_lookup.json from the ecoinvent properties export.

    python calculators/agri_diesel/build_property_defaults.py

The export (PROPERTIES_EXPORT_FILE in config_new_datasets.py, ~230 MB) has one
row per property of every exchange of every ecoinvent dataset.  What the
pipeline needs is much smaller: for each intermediate product, the value each
of its properties usually has.  So, per (exchange name, exchange unit), every
property is reduced to its most common amount, with how many datasets carry
it and how many of those agree -- a product whose datasets disagree is still
listed, but the agreement says how far to trust the value.

Re-run after an ecoinvent release; the pipeline never reads the export itself.

Author: Andreas Giakoumatos
"""

import os
import sys
import json
import collections
from datetime import date

# The project root -- the folder holding config_new_datasets.py -- on the
# path, wherever in the project this file sits.
_ROOT = os.path.dirname(os.path.abspath(__file__))
while not os.path.isfile(os.path.join(_ROOT, 'config_new_datasets.py')):
    _ROOT = os.path.dirname(_ROOT)
sys.path.insert(0, _ROOT)
from config_new_datasets import PROPERTIES_EXPORT_FILE, PROPERTY_DEFAULTS_FILE

WANTED = ('exchange name', 'exchange unitName', 'group', 'property name',
          'property unitName', 'property amount', 'property mathematicalRelation')


def main() -> int:
    from openpyxl import load_workbook
    print(f'Reading {PROPERTIES_EXPORT_FILE} ...')
    wb = load_workbook(PROPERTIES_EXPORT_FILE, read_only=True, data_only=True)
    values: dict = collections.defaultdict(collections.Counter)  # (name, unit, prop, punit) -> amounts
    rows = 0
    for ws in wb.worksheets:
        header = None
        for row in ws.iter_rows(values_only=True):
            if header is None:
                header = {str(h).strip(): i for i, h in enumerate(row) if h is not None}
                missing = [w for w in WANTED if w not in header]
                if missing:
                    print(f'  sheet {ws.title!r}: columns missing {missing} -- skipped')
                    break
                continue
            rows += 1
            if rows % 500000 == 0:
                print(f'  {rows:,} rows ...', flush=True)
            group = str(row[header['group']] or '')
            if 'Environment' in group:
                continue                     # elementary flows carry no product properties
            name = str(row[header['exchange name']] or '').strip()
            prop = str(row[header['property name']] or '').strip()
            amount = row[header['property amount']]
            if not name or not prop or not isinstance(amount, (int, float)):
                continue
            key = (name, str(row[header['exchange unitName']] or '').strip(), prop,
                   str(row[header['property unitName']] or '').strip())
            values[key][round(float(amount), 10)] += 1
    wb.close()

    lookup: dict = {}
    for (name, unit, prop, punit), counter in values.items():
        amount, agree = counter.most_common(1)[0]
        entry = lookup.setdefault(f'{name}|{unit}', {})
        entry[prop] = {'amount': amount, 'unitName': punit,
                       'datasets': sum(counter.values()), 'agree': agree}
    out = {'_source': os.path.basename(PROPERTIES_EXPORT_FILE),
           '_built': date.today().isoformat(), 'products': lookup}
    with open(PROPERTY_DEFAULTS_FILE, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=0, sort_keys=True)
    print(f'{rows:,} rows read; {len(lookup):,} products written to {PROPERTY_DEFAULTS_FILE}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
